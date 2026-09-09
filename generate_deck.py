from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Color Palette (Deep Tech Theme)
BG_DARK = RGBColor(15, 23, 42)      # Slate Dark
CARD_BG = RGBColor(30, 41, 59)      # Slate Card
TEXT_LIGHT = RGBColor(248, 250, 252) # White
ACCENT_BLUE = RGBColor(56, 189, 248) # Neon Cyan
ACCENT_PURPLE = RGBColor(168, 85, 247) # Neon Purple
TEXT_MUTED = RGBColor(148, 163, 184) # Gray

def set_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

def add_header(slide, title_text, category_text="VOICEPULSE AI // SYSTEM ARCHITECTURE"):
    # Category Tracker
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
    tf_cat = cat_box.text_frame
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_PURPLE
    
    # Action Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_LIGHT

# ==========================================
# SLIDE 1: Title Slide (Dark Immersive)
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_background(slide1)

title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0))
tf = title_box.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "VoicePulse AI"
p1.font.size = Pt(54)
p1.font.bold = True
p1.font.color.rgb = ACCENT_BLUE

p2 = tf.add_paragraph()
p2.text = "Real-Time Voice QA & Enterprise Speech Analytics Platform"
p2.font.size = Pt(22)
p2.font.color.rgb = TEXT_LIGHT
p2.space_before = Pt(10)

p3 = tf.add_paragraph()
p3.text = "Production-Ready Microservice Architecture • Built with FastAPI, Docker & AssemblyAI"
p3.font.size = Pt(14)
p3.font.color.rgb = TEXT_MUTED
p3.space_before = Pt(25)

# ==========================================
# SLIDE 2: Executive Summary / The Problem
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_background(slide2)
add_header(slide2, "The Challenge: Manual QA is Slow & Unscalable")

cards_data = [
    ("The Latency Bottleneck", "Traditional speech review processes rely on manual spot-checking, causing severe evaluation delays across high-volume customer interaction lines."),
    ("Compliance Blind Spots", "Human oversight misses up to 85% of regulatory or script adherence infractions during live agent-customer audio calls."),
    ("Integration Friction", "Existing analytics tools lack lightweight, plug-and-play containerized deployment frameworks required by modern cloud infrastructures.")
]

for i, (head, desc) in enumerate(cards_data):
    left = Inches(0.8 + (i * 3.9))
    shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.2), Inches(3.6), Inches(4.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = ACCENT_PURPLE
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.4)
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    
    p = tf.paragraphs[0]
    p.text = head
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    p_body = tf.add_paragraph()
    p_body.text = desc
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_MUTED
    p_body.space_before = Pt(15)

# ==========================================
# SLIDE 3: System Workflow & Architecture
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_background(slide3)
add_header(slide3, "Working Flow Diagram: End-to-End Pipeline")

steps = [
    ("1. Ingestion Layer", "Audio input streams or payloads sent securely to FastAPI REST endpoints."),
    ("2. Auth & Processing", "Token validation, schema validation with Pydantic v2, and queue sorting."),
    ("3. Speech Engine", "AssemblyAI streaming integration processes audio transcripts instantly."),
    ("4. QA Evaluation", "Automated scoring, sentiment calculation, and database persistence.")
]

for i, (title, text) in enumerate(steps):
    top = Inches(2.0 + (i * 1.2))
    
    # Step Node Box
    shape = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(3.2), Inches(0.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = ACCENT_BLUE
    
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    # Description Box
    desc_box = slide3.shapes.add_textbox(Inches(4.3), top, Inches(8.0), Inches(0.9))
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = text
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = TEXT_LIGHT

# ==========================================
# SLIDE 4: Technical Stack & Production Readiness
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_background(slide4)
add_header(slide4, "Enterprise Technology Stack & Docker Deployment")

tech_columns = [
    ("Core Backend", ["FastAPI (Asynchronous)", "Uvicorn ASGI Server", "Pydantic v2 Validation", "Python 3.13 Runtime"]),
    ("Data & Integration", ["AssemblyAI Speech API", "SQLAlchemy ORM", "SQLite Persistence", "Requests Network Layer"]),
    ("DevOps & Production", ["Docker Containerization", "Multi-stage Dockerfile", "Docker Compose Config", "Automated Health Checks"])
]

for i, (col_title, items) in enumerate(tech_columns):
    left = Inches(0.8 + (i * 3.9))
    shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.2), Inches(3.6), Inches(4.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = ACCENT_BLUE
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.4)
    tf.margin_left = Inches(0.3)
    
    p = tf.paragraphs[0]
    p.text = col_title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    
    for item in items:
        pi = tf.add_paragraph()
        pi.text = f"• {item}"
        pi.font.size = Pt(14)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(12)

# Save Presentation
prs.save("VoicePulse_AI_Presentation.pptx")
print("Presentation generated successfully: VoicePulse_AI_Presentation.pptx")